"""Put the real video clip back into a PPTX downloaded from the deck.

The deck's download shows only the still frame of a clip. This swaps that
picture for the .mp4 at the same place and size, so the clip plays in PowerPoint.

Usage:  python embed_clips.py "WiEvac.pptx"
Writes: WiEvac_co_video.pptx next to the input.
"""
import sys
import tempfile
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
# slide number (1-based, as shown in the deck) -> clip file in this folder
CLIPS = {3: "clip_demo_app_2_dien_thoai.mp4", 7: "clip_dinh_vi.mp4", 12: "clip_csi_nguoi_di_qua.mp4"}


def largest_picture(slide):
    pictures = [s for s in slide.shapes if s.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
    return max(pictures, key=lambda s: s.width * s.height) if pictures else None


def main(path: Path) -> None:
    prs = Presentation(str(path))
    for number, clip in CLIPS.items():
        slide = prs.slides[number - 1]
        still = largest_picture(slide)
        if still is None:
            print(f"slide {number}: no picture found, skipped")
            continue
        with tempfile.NamedTemporaryFile(suffix="." + still.image.ext, delete=False) as poster:
            poster.write(still.image.blob)
        left, top, width, height = still.left, still.top, still.width, still.height
        still._element.getparent().remove(still._element)
        slide.shapes.add_movie(str(HERE / clip), left, top, width, height,
                               poster_frame_image=poster.name, mime_type="video/mp4")
        print(f"slide {number}: clip {clip} embedded")
    out = path.with_name(path.stem + "_co_video.pptx")
    prs.save(str(out))
    print("saved", out)


if __name__ == "__main__":
    main(Path(sys.argv[1]))

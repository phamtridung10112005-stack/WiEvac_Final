# -*- coding: utf-8 -*-
"""Báo cáo vòng bán kết VEDC 2026 — WiEvac - UTEHY. Tối đa 10 trang."""
import io
import sys
from pathlib import Path

import docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parent
NAVY = RGBColor(16, 54, 112)
FONT = "Times New Roman"


def shade(cell, color):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    )


def margins(cell, top=60, bottom=60, left=80, right=80):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f"</w:tcMar>"
        )
    )


def borders(table, color="B0C4DE", sz="4"):
    table._tbl.tblPr.append(
        parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f"</w:tblBorders>"
        )
    )


def no_borders(table):
    table._tbl.tblPr.append(
        parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            + "".join(
                f'<w:{edge} w:val="nil" w:sz="0" w:space="0" w:color="auto"/>'
                for edge in ("top", "bottom", "left", "right", "insideH", "insideV")
            )
            + "</w:tblBorders>"
        )
    )


def signature_stream():
    """Crop the scanned signature so the white margin does not eat a page."""
    from PIL import Image

    image = Image.open(ROOT / "chuky.jpg").convert("RGB")
    mask = image.convert("L").point(lambda p: 255 if p < 245 else 0)
    box = mask.getbbox()
    if box:
        pad = 12
        left = max(0, box[0] - pad)
        top = max(0, box[1] - pad)
        right = min(image.width, box[2] + pad)
        bottom = min(image.height, box[3] + pad)
        image = image.crop((left, top, right, bottom))
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    buf.seek(0)
    return buf


def font_run(run, size=13, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = FONT
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = parse_xml(
            f'<w:rFonts {nsdecls("w")} w:ascii="{FONT}" w:hAnsi="{FONT}" w:eastAsia="{FONT}" w:cs="{FONT}"/>'
        )
        rpr.append(rfonts)
    else:
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), FONT)


def build():
    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(1.6)
        section.header_distance = Cm(0.6)
        section.footer_distance = Cm(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(13)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.08
    pf.space_after = Pt(2)
    pf.space_before = Pt(0)

    def P(text="", *, bold=False, italic=False, size=13, center=False, before=0, after=2, color=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.08
        if text:
            font_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
        return p

    def mixed(parts, *, center=False, before=0, after=2, size=13):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.08
        for text, flags in parts:
            font_run(
                p.add_run(text),
                size=size,
                bold=flags.get("b", False),
                italic=flags.get("i", False),
                color=flags.get("c"),
            )
        return p

    def H(text):
        return P(text, bold=True, size=13, before=8, after=2, color=NAVY)

    def H2(text):
        return P(text, bold=True, size=13, before=6, after=2)

    def bullet(prefix, text):
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.08
        if prefix:
            font_run(p.add_run(prefix), size=13, bold=True)
        font_run(p.add_run(text), size=13)
        return p

    def table(headers, rows, widths):
        tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        borders(tbl)
        for i, h in enumerate(headers):
            cell = tbl.rows[0].cells[i]
            cell.text = ""
            cell.width = widths[i]
            shade(cell, "E7EEF6")
            margins(cell)
            pp = cell.paragraphs[0]
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font_run(pp.add_run(h), size=13, bold=True)
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = tbl.rows[r + 1].cells[c]
                cell.text = ""
                cell.width = widths[c]
                if r % 2 == 1:
                    shade(cell, "F7FAFC")
                margins(cell)
                pp = cell.paragraphs[0]
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER if c == 0 else WD_ALIGN_PARAGRAPH.LEFT
                font_run(pp.add_run(val), size=13)
        return tbl

    P("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, center=True, after=0)
    P("Độc lập - Tự do - Hạnh phúc", bold=True, center=True, after=0)
    P("---------------", center=True, after=6)
    P(
        "BÁO CÁO DỰ ÁN VÒNG 2 (BÁN KẾT)\nCUỘC THI “THIẾT KẾ ĐIỆN TỬ VIỆT NAM 2026”",
        bold=True,
        size=14,
        center=True,
        after=2,
        color=NAVY,
    )
    P("Chủ đề: AI nhúng cho hệ thống IoT thông minh", bold=True, italic=True, size=13, center=True, after=0)
    P("Lưu ý: Nội dung báo cáo không quá 10 trang", italic=True, size=13, center=True, after=6, color=RGBColor(90, 90, 90))

    H("1. Thông tin chung")
    mixed([("Tên đội thi: ", {"b": True}), ("WiEvac - UTEHY", {})])
    mixed([
        ("Trường/Viện/Đơn vị đào tạo: ", {"b": True}),
        ("Trường Đại học Công nghệ Kỹ thuật Hưng Yên", {}),
    ])

    H("2. Thông tin dự án")
    H2("2.1. Tên dự án")
    P("WiEvac – Hệ thống điều phối sơ tán khẩn cấp thông minh ứng dụng sóng WiFi CSI.")

    H2("2.2. Lĩnh vực ứng dụng")
    P(
        "☐ Công nghiệp          ☐ Nông nghiệp          ☐ Y tế\n"
        "☐ Giáo dục             ☐ Giao thông           ☑ Đô thị thông minh\n"
        "☐ Khác: ........................",
        after=2,
    )

    H2("2.3. Ý tưởng chính")
    P(
        "Ở nhà trọ, chung cư mini và nhiều công trình trong đô thị, khói làm mất tầm nhìn rất nhanh. "
        "Biển chỉ lối thoát gắn cố định không biết phía trước còn đi được hay không. Camera gần như hết tác dụng khi ống kính bị khói che, "
        "và cũng khó lắp sát cửa phòng vì đụng đến chuyện riêng tư. Đầu báo khói cho biết có sự cố, nhưng không nói đoạn hành lang nào còn thông."
    )
    P(
        "WiEvac đặt các cặp ESP32-S3 dọc hành lang. Một node phát sóng Wi-Fi, node kia thu và xem sóng bị suy hao, phản xạ và dao động thế nào trên đúng đoạn đó. "
        "Từ mức suy hao và mức dao động so với lúc hành lang đang trống, node ước lượng đoạn còn dễ đi hay đang bị cản. "
        "Kết quả gửi về Raspberry Pi 5 để theo dõi các đoạn và phát lệnh chỉ đường. Trên node, DFPlayer Mini phát câu hướng dẫn, "
        "PAM8403 đưa tiếng ra loa, matrix LED hiện ký hiệu cùng hướng."
    )
    P(
        "Kèm theo là ứng dụng trên điện thoại. Ứng dụng quét hiện trường lắp đặt và dựng lại mô hình không gian ba chiều, "
        "để nhìn được chỗ đặt node trong tòa nhà. Khi cần sơ tán, hệ thống gửi thông báo về máy đã cài ứng dụng. "
        "Với người đang ở trong công trình, ứng dụng còn dùng để bám vị trí của họ, phục vụ việc điều hướng ra ngoài."
    )
    P(
        "Cách đo trên node bám tình trạng sóng của chính đoạn đang lắp, không nhận diện khuôn mặt hay đếm người bằng camera. "
        "Khi sóng nhiễu hoặc dữ liệu chưa ổn, hệ thống khoan kết luận, tránh chỉ một hướng thật chắc trong lúc chưa đủ căn cứ."
    )

    H("3. Mô tả chi tiết")
    H2("3.1. Mục tiêu và vấn đề thực tiễn cần giải quyết")
    P(
        "Nhiều công trình dân sinh có hành lang hẹp, ít lối ra, cầu thang lại dễ thành đường khói đi lên. "
        "Khi sự cố xảy ra, người bên trong phải chọn hướng rất nhanh trong lúc mắt gần như không còn nhìn thấy. "
        "Đèn thoát hiểm chỉ một chiều cố định, nên có thể dẫn vào đúng phía đang tắc. "
        "Radar chuyên dụng cảm nhận được xuyên khói, nhưng giá và cách lắp khó phù hợp nhà trọ hay chung cư mini."
    )
    P(
        "WiEvac nhắm một lớp hỗ trợ chỉ đường: dùng sóng Wi-Fi giữa các node để ước lượng đoạn hành lang còn dễ đi hay đang bị cản, "
        "rồi báo ngay tại chỗ bằng giọng nói và ký hiệu đèn. Kết quả dùng để chọn hướng, bám độ thông thoáng của đoạn. "
        "Cách bố trí theo cụm node nên có thể thêm đoạn, thêm tầng mà không phải đổi cả nguyên lý đo."
    )

    H2("3.2. Phương pháp tiếp cận và giải pháp công nghệ")
    P(
        "Sóng Wi-Fi giữa node phát và node thu không đi một đường thẳng. Sóng đập tường, cửa, người và đồ vật rồi cộng lại ở đầu thu. "
        "Khi có người đi qua hoặc đứng chắn, độ suy hao tăng và sóng dao động mạnh hơn lúc hành lang trống. "
        "Node thu chỉ lấy sóng từ đúng node phát đã ghép, bỏ các khung nhiễu hoặc khung hỏng, rồi so với đường nền của chính đoạn đó."
    )
    P(
        "Đường nền là mức sóng lúc đoạn đang yên. Đường này được học tại chỗ, trên ESP32-S3, nên đổi hành lang hoặc đổi chỗ gắn anten thì hệ thống làm quen lại, "
        "không mang ngưỡng của một phòng cũ sang phòng mới. Đoạn được xem là thông khi sóng gần với đường nền. "
        "Đoạn được xem là kém hoặc tắc khi độ suy hao và mức dao động cùng lệch rõ, và lệch đó kéo dài chứ không chỉ một nhịp nhiễu ngắn. "
        "Nếu sóng yếu vì mất liên lạc, vì nhiễu dày hoặc vì mới bật máy chưa ổn, node chưa kết luận, tránh báo tắc hoặc báo thông một cách chắc chắn."
    )
    P(
        "Mỗi đoạn gửi về Raspberry Pi 5 một kết quả gọn: mức thông thoáng, tình trạng đoạn và độ tin cậy của lần đo đó. "
        "Máy chủ kiểm tra gói tin, theo dõi từng đoạn riêng và lưu lại quá trình chạy. Trên màn hình theo dõi có thể xem đoạn nào còn liên lạc, đoạn nào đang yếu, "
        "rồi phát lệnh chỉ dẫn xuống một node hoặc cả cụm. Node nhận lệnh sẽ nói hướng đi, hướng rẽ, hướng quay lại, cửa thoát hoặc yêu cầu dừng, đồng thời hiện ký hiệu trên matrix LED."
    )
    P(
        "Phần nặng nằm trên node để khỏi phụ thuộc đường truyền ra ngoài tòa nhà. Máy chủ giữ vai trò gom các đoạn và điều hướng chỉ dẫn. "
        "Khi mặt bằng rộng hơn, có thể nối các đoạn thành một mạch đi cho cả tầng mà không phải sửa cách đo của từng node."
    )

    H2("3.3. Mô hình hệ thống / kiến trúc thiết kế")
    P("Hệ thống có bốn khối. Phần ước lượng độ thông thoáng nằm trên node thu. Máy chủ gom các đoạn và phát lệnh. Điện thoại là lối chỉ dẫn tới từng người đã cài ứng dụng.")
    bullet("Khối đo. ", "Cặp ESP32-S3. Node phát duy trì sóng đo. Node thu theo dõi độ suy hao và mức dao động, so với đường nền của đoạn, rồi đưa ra mức thông thoáng.")
    bullet("Khối tổng hợp. ", "Raspberry Pi 5 nhận kết quả từng đoạn, lưu diễn biến và hiện trên màn hình theo dõi. Từ đây hệ thống chọn hướng sơ tán và phát lệnh chỉ dẫn.")
    bullet("Khối chỉ đường tại node. ", "DFPlayer Mini phát lời hướng dẫn, PAM8403 khuếch đại ra loa, matrix LED hiện ký hiệu cùng hướng để người vừa nghe vừa nhìn.")
    bullet("Khối trên điện thoại. ", "Ứng dụng quét khu vực lắp đặt và dựng mô hình không gian ba chiều của hiện trường. Ứng dụng gửi thông báo sơ tán về máy khi cần rời khỏi tòa nhà, đồng thời theo vị trí người đã cài ứng dụng trong công trình để việc điều hướng bám được người đang ở bên trong.")
    P(
        "Nếu một node không về thẳng được máy chủ, node bên cạnh có thể chuyển tiếp giúp. "
        "Trên màn hình vẫn thấy đoạn đó còn đường liên lạc. Khi mất hẳn, đoạn được báo mất kết nối, không dùng số liệu cũ để chỉ một hướng mới."
    )
    P("Luồng một lần đo:", before=4, after=2)
    table(
        ["Bước", "Khối làm", "Đầu ra"],
        [
            ("1", "ESP32-S3 phát và thu", "Sóng Wi-Fi trên đoạn hành lang giữa đúng cặp node"),
            ("2", "Chương trình trên node thu", "Mức thông thoáng của đoạn và độ tin cậy của lần đo"),
            ("3", "Raspberry Pi 5", "Tình trạng các đoạn, lịch sử chạy, lệnh chỉ dẫn"),
            ("4", "Loa và matrix LED", "Câu hướng dẫn và ký hiệu tại chỗ"),
            ("5", "Ứng dụng điện thoại", "Mô hình 3D hiện trường, thông báo sơ tán, vị trí người đã cài ứng dụng"),
        ],
        [Cm(1.6), Cm(5.0), Cm(9.6)],
    )

    H2("3.4. Công cụ, linh kiện và công cụ thiết kế")
    P(
        "Mạch node được vẽ bằng KiCad. Hình 1 là phần bố trí mạch in: cụm chân cắm module, điện trở, tụ và nút nhấn trên một mạch gọn, đủ gắn các khối của node. "
        "Chương trình trong ESP32-S3, phần mềm trên Raspberry Pi 5 và ứng dụng điện thoại được viết bằng Visual Studio Code."
    )
    P("Linh kiện chính:", before=4, after=2)
    table(
        ["TT", "Linh kiện", "Vai trò trong hệ thống"],
        [
            ("1", "ESP32-S3", "Phát hoặc thu sóng, ước lượng độ thông thoáng của đoạn, nhận lệnh chỉ dẫn"),
            ("2", "Raspberry Pi 5", "Gom kết quả các đoạn, màn hình theo dõi, phát lệnh chỉ đường"),
            ("3", "DFPlayer Mini", "Phát câu hướng dẫn bằng giọng nói theo lệnh"),
            ("4", "PAM8403", "Khuếch đại tiếng từ DFPlayer Mini ra loa trên node"),
            ("5", "Matrix LED", "Hiện ký hiệu hướng, đi cùng nội dung phát trên loa"),
        ],
        [Cm(1.4), Cm(3.8), Cm(11.0)],
    )

    H("4. Kết quả đạt được đến vòng bán kết")
    H2("4.1. Mô hình / sản phẩm thử nghiệm")
    P(
        "Đến vòng bán kết, nhóm đã có mẫu node như Hình 2: vỏ đứng, ăng-ten ngoài và loa phía trước. "
        "Mạch gắn trong node được bố trí trên KiCad, thể hiện ở Hình 1. "
        "Phần chạy được gồm cặp node đo sóng trên một đoạn, node thu ước lượng độ thông thoáng rồi gửi về Raspberry Pi 5. "
        "Màn hình theo dõi hiện tình trạng từng đoạn và gửi lệnh để loa trên node phát câu hướng dẫn. "
        "Trên điện thoại, ứng dụng đã quét và dựng được mô hình không gian ba chiều của hiện trường lắp đặt."
    )

    H2("4.2. Các tính năng đã hoàn thiện")
    bullet("Ước lượng từng đoạn. ", "Node thu theo dõi độ suy hao và mức dao động của sóng, đối chiếu với đường nền lúc đoạn đang trống, rồi đưa ra mức thông thoáng của đúng đoạn đó. Đổi chỗ lắp thì đường nền được học lại.")
    bullet("Lọc nhiễu trước khi kết luận. ", "Khung sóng hỏng, mất liên lạc hoặc dao động quá ngắn không bị hiểu thành có vật cản. Khi dữ liệu chưa đủ tin, node chưa chốt đoạn thông hay tắc.")
    bullet("Theo dõi tại máy chủ. ", "Raspberry Pi 5 nhận kết quả đã tính, lưu diễn biến và hiện mức thông thoáng cùng tình trạng liên lạc của từng đoạn.")
    bullet("Chỉ dẫn bằng âm thanh. ", "Lệnh đi thẳng, rẽ, quay lại, hướng tới cửa thoát hoặc yêu cầu dừng được DFPlayer Mini phát thành lời, qua PAM8403 ra loa.")
    bullet("Chuyển tiếp khi mất đường. ", "Node không về thẳng được máy chủ thì có thể nhờ node còn liên lạc chuyển tiếp. Mất hẳn thì màn hình báo mất kết nối.")
    bullet("Mô hình hiện trường trên điện thoại. ", "Ứng dụng quét khu vực lắp đặt và dựng lại mô hình không gian ba chiều, dùng để xem bố trí node trong tòa nhà. Khi sơ tán, thông báo được gửi về các máy đã cài ứng dụng.")

    H2("4.3. Thông số kỹ thuật / Benchmark (nếu có)")
    P("")
    P("")

    H2("4.4. Hình ảnh minh họa")
    pics = doc.add_table(rows=2, cols=2)
    pics.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(pics)
    left, right = pics.rows[0].cells
    cap_l, cap_r = pics.rows[1].cells
    for cell in (left, right, cap_l, cap_r):
        margins(cell, top=40, bottom=40, left=60, right=60)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    left.paragraphs[0].add_run().add_picture(str(ROOT / "mach.jpg"), width=Inches(4.15))
    right.paragraphs[0].add_run().add_picture(str(ROOT / "node.jpg"), width=Inches(1.85))
    font_run(
        cap_l.paragraphs[0].add_run("Hình 1. Bố trí mạch node trên KiCad."),
        size=13,
        italic=True,
    )
    font_run(
        cap_r.paragraphs[0].add_run("Hình 2. Mẫu node: vỏ, ăng-ten và loa."),
        size=13,
        italic=True,
    )

    H("5. Kế hoạch hoàn thiện cho vòng chung kết")
    H2("5.1. Hạng mục cần bổ sung / tối ưu")
    bullet("Đường nền ổn hơn theo thời gian. ", "Theo dõi thêm ảnh hưởng của nhiệt độ và thời gian máy chạy, để mức sóng trôi chậm không bị đọc thành có người hoặc làm mất người đang đứng.")
    bullet("Chỉ dẫn cho cả mặt bằng. ", "Nối tình trạng nhiều đoạn thành một hướng sơ tán cho hành lang và cho một tầng. Loa, matrix LED và thông báo trên điện thoại bám cùng hướng đó.")
    bullet("Vị trí người trong tòa nhà. ", "Siết phần định vị của người đã cài ứng dụng, để lệnh sơ tán bám được người đang đứng ở khu nào của công trình.")
    bullet("Lớp đối chiếu khi đã có dữ liệu. ", "Sau khi thu đủ các kiểu hành lang, có thể thêm một lớp học máy trên node để đối chiếu mức thông thoáng. Lớp này hỗ trợ kết luận, không tự quyết khi sóng đang nhiễu.")
    bullet("Chạy trong điều kiện khó hơn. ", "Đo trên hành lang dài hơn, thêm tình huống khói mô phỏng an toàn, và tính nguồn dự phòng để node còn nói được khi mất điện lưới.")

    H2("5.2. Kế hoạch triển khai tiếp theo")
    table(
        ["Giai đoạn", "Việc chính", "Đầu ra"],
        [
            ("1", "Chốt mạch node, loa, matrix LED và anten", "Cụm node mang đi đo trên hành lang"),
            ("2", "Thu số liệu lúc trống và lúc có người, theo dõi trôi nền theo nhiệt", "Đường nền ổn hơn trên đoạn đang lắp"),
            ("3", "Nối các đoạn thành hướng sơ tán; gắn với thông báo trên điện thoại", "Loa, đèn và điện thoại cùng một hướng"),
            ("4", "Siết định vị người đã cài ứng dụng; chạy thử khói mô phỏng", "Buổi trình diễn cho vòng chung kết"),
        ],
        [Cm(2.4), Cm(8.2), Cm(5.6)],
    )

    H2("5.3. Phân công nhân sự trong nhóm")
    table(
        ["Họ và tên", "Phần việc phụ trách"],
        [
            ("Phạm Trí Dũng", "Đo sóng trên từng node và ước lượng mức thông thoáng của từng đoạn hành lang."),
            ("Chử Quốc Hào", "Phần cứng node, cùng phần định vị người đã cài ứng dụng trong tòa nhà."),
            ("Nguyễn Phú Anh", "Xử lý sơ tán và điều hướng: ghép tình trạng các đoạn thành hướng đi, ra lệnh cho loa, đèn và thông báo trên điện thoại."),
            ("Nguyễn Tuấn Kiệt", "Ứng dụng điện thoại, gồm quét hiện trường và dựng mô hình không gian ba chiều của tòa nhà."),
        ],
        [Cm(4.2), Cm(12.0)],
    )

    H("6. Đánh giá và đề xuất")
    H2("6.1. Điểm mạnh của dự án")
    bullet("Còn dùng được khi mất tầm nhìn. ", "Sóng Wi-Fi 2,4 GHz không phụ thuộc ống kính. Khói che camera thì đường sóng giữa hai node vẫn còn để ước lượng đoạn có bị cản hay không.")
    bullet("Xử lý ngay trên node. ", "Mỗi đoạn tự so độ suy hao và mức dao động với đường nền của mình. Máy chủ nhận kết quả đã rút gọn, nên chỉ dẫn không phải chờ đưa toàn bộ sóng ra ngoài tòa nhà.")
    bullet("Đổi hành lang thì học lại. ", "Đường nền lấy tại chỗ đang lắp. Không dùng một ngưỡng sóng cố định mang từ phòng này sang phòng khác.")
    bullet("Chưa đủ tin thì chưa chỉ chắc. ", "Nhiễu, mất liên lạc hoặc mới khởi động không bị biến thành một hướng thoát chắc chắn. Chỉ dẫn bám những đoạn còn số liệu mới.")
    bullet("Không cần camera. ", "Hành lang nhà ở không phải gắn camera sát cửa phòng. Người trong khói nghe loa và nhìn ký hiệu trên node.")
    bullet("Linh kiện dễ kiếm. ", "ESP32-S3, DFPlayer Mini, PAM8403, matrix LED và một Raspberry Pi 5. Có thể thêm node khi mặt bằng dài hơn.")

    H2("6.2. Hạn chế và thách thức")
    bullet(
        "Trôi đường nền do nhiệt. ",
        "Nhiệt độ phòng, độ ẩm và bản thân mạch chạy lâu đều làm mức sóng trôi chậm, dù hành lang không có thêm người. "
        "Nếu giữ mãi đường nền lúc mới bật máy, sau một lúc hệ thống dễ tưởng đoạn đang có người, hoặc ngược lại bỏ qua một thay đổi thật. "
        "Hướng xử lý: tách biến đổi chậm của môi trường khỏi biến đổi nhanh khi có người đi qua. "
        "Khi đoạn yên, đường nền được làm mới theo chính đoạn đó, không khóa một mốc đo từ lúc lắp. "
        "Phần này tiếp tục được theo dõi khi đo lâu hơn trên hành lang, vì mức trôi còn phụ thuộc chỗ đặt node và thời tiết trong nhà.",
    )
    bullet(
        "Chưa đưa mô hình nhận diện vào kết luận. ",
        "Đếm người, nhận mặt hay phân loại vật cản bằng một mạng học sẵn cần rất nhiều mẫu đã gán nhãn, lại dễ sai khi đổi hành lang, đổi hướng anten hoặc khi có khói. "
        "ESP32-S3 cũng không phải chỗ để nhét một mạng lớn nếu vẫn muốn phản hồi kịp lúc sự cố. "
        "Vì vậy giai đoạn này nhóm không dùng nhận diện để quyết định đoạn tắc hay thông. "
        "Kết luận dựa trên độ suy hao và mức dao động của sóng so với đường nền tại chỗ: chỉ cần biết đoạn còn dễ đi hay đang bị cản, không cần biết đó là ai. "
        "Khi đã có thêm số liệu từng kiểu hành lang, nhóm có thể bổ sung một lớp học máy nhỏ để đối chiếu. "
        "Lớp đó chỉ hỗ trợ. Nếu sóng đang nhiễu hoặc lớp đối chiếu lệch với mức suy hao, kết luận không bị kéo theo phía đẹp hơn.",
    )
    bullet(
        "Cùng một người nhưng sóng đổi theo chỗ đứng. ",
        "Hành lang hẹp, cửa sắt và tường làm sóng phản xạ khác nhau dù người đi cùng một kiểu. "
        "Hướng xử lý: so với đường nền của đúng đoạn đang đo, và chỉ nâng mức cảnh báo khi độ suy hao cùng mức dao động kéo dài. "
        "Anten ngoài trên node để chỉnh hướng sóng dọc hành lang lúc lắp. "
        "Các mặt bằng phức tạp hơn sẽ được nối dần từ kết quả từng đoạn, không gán một quy tắc chung cho mọi tòa nhà.",
    )

    H2("6.3. Đề xuất hỗ trợ từ BTC (thiết bị, tài liệu, cố vấn,...)")
    P("")
    P("")
    P("", after=6)
    sign = doc.add_table(rows=1, cols=2)
    sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(sign)
    sign.rows[0].cells[0].text = ""
    cell = sign.rows[0].cells[1]
    cell.text = ""
    lines = [
        ("Hưng Yên, ngày …… tháng …… năm 2026", True, False),
        ("ĐẠI DIỆN ĐỘI THI", False, True),
        ("(Ký và ghi rõ họ tên)", True, False),
    ]
    for i, (text, italic, bold) in enumerate(lines):
        pp = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pp.paragraph_format.space_after = Pt(0)
        pp.paragraph_format.line_spacing = 1.0
        font_run(pp.add_run(text), size=13, bold=bold, italic=italic)
    sig = cell.add_paragraph()
    sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig.paragraph_format.space_before = Pt(2)
    sig.paragraph_format.space_after = Pt(0)
    sig.add_run().add_picture(signature_stream(), width=Cm(4.2))
    name = cell.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_before = Pt(0)
    name.paragraph_format.space_after = Pt(0)
    font_run(name.add_run("Phạm Trí Dũng"), size=13, bold=True)

    out = ROOT / "3._bao_cao_du_an_vong_ban_ket_cuoc_thi_vedc_2026_ten_doi.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    build()
